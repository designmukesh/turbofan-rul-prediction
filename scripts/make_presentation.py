"""Build presentation/turbofan_rul_presentation.pptx — 12 slides, real charts only."""
import sys
from pathlib import Path

sys.path.insert(0, "..")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "reports" / "figures"
OUT = ROOT / "presentation" / "turbofan_rul_presentation.pptx"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
LEFT = Inches(0.7)
CONTENT_W = Inches(11.9)


def add_slide(title, bullets=None, img=None, img_w=Inches(6), note=None):
    s = prs.slides.add_slide(BLANK)
    # title bar
    tx = s.shapes.add_textbox(LEFT, Inches(0.25), CONTENT_W, Inches(0.8)).text_frame
    tx.word_wrap = True
    p = tx.paragraphs[0]
    p.text, p.font.size, p.font.bold = title, Pt(30), True
    p.font.name = "Calibri"
    top = Inches(1.3)
    if img:
        s.shapes.add_picture(str(img), LEFT, top, width=img_w)
        left2 = LEFT + img_w + Inches(0.5)
        w2 = CONTENT_W - img_w - Inches(0.5)
    else:
        left2, w2 = LEFT, CONTENT_W
    if bullets:
        tf = s.shapes.add_textbox(left2, top, w2, Inches(5.5)).text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.text, par.font.size = b, Pt(19)
            par.font.name = "Calibri"
            par.space_after, par.level = Pt(10), 0
    if note:
        tf = s.shapes.add_textbox(LEFT, Inches(6.7), CONTENT_W, Inches(0.6)).text_frame
        tf.word_wrap = True
        par = tf.paragraphs[0]
        par.text, par.font.size = note, Pt(14)
        par.font.italic = True
        par.font.name = "Calibri"
    return s


add_slide("Turbofan Engine Remaining-Useful-Life Prediction",
          ["NASA C-MAPSS FD001 — 100 run-to-failure engines, 21 sensors",
           "Question: how many cycles until this engine fails?",
           "Answer: HistGradientBoosting, test RMSE 12.74 cycles",
           "Honest band from repeated CV: ~14–15 cycles",
           "Independent portfolio project — public data only"],
          note="Mukesh Sharma · github.com/designmukesh/turbofan-rul-prediction")

add_slide("The business problem",
          ["Unplanned engine removals cost 3–10× a planned shop visit",
           "Fixed-interval overhauls waste healthy life; surprises risk failure",
           "The decision to improve: WHEN to pull an engine",
           "Users: reliability engineers, maintenance planners, fleet managers",
           "Late predictions (saying 'healthy' when it's not) are the dangerous error"])

add_slide("The dataset — and what it isn't",
          ["NASA C-MAPSS FD001 (Saxena & Goebel, 2008), PHM Society mirror",
           "100 train engines (20,631 cycles) + 100 truncated test engines",
           "21 sensors + 3 operating settings per cycle",
           "SIMULATED data — thermodynamic model, not field data",
           "One operating condition, one fault mode (HPC degradation)",
           "Every claim in this deck is conditional on that"])

add_slide("EDA: the signal is visible before any model",
          ["7 of 21 sensors never move → dropped; 14 carry degradation trend",
           "Temperatures rise, speeds shift in the last ~100 cycles",
           "Lifespans vary 128–362 cycles — age alone predicts nothing"],
          img=FIG / "02_sensor_trends.png", img_w=Inches(7.2))

add_slide("Solution architecture",
          ["per-cycle sensors → 30-cycle sliding windows",
           "→ 84 statistical features (mean/std/min/max/last/slope × 14 sensors)",
           "→ HistGradientBoosting → predicted RUL + trend",
           "→ risk band → illustrative action → HUMAN review → work order",
           "→ outcome feedback → retrain",
           "The model advises; certified processes decide."],
          note="Batch scoring daily is enough — RUL moves slowly.")

add_slide("Feature engineering: transparent by design",
          ["One cycle has no trend — windows give the model history",
           "Six statistics per sensor: distribution (mean/std/min/max),",
           "recency (last), and within-window slope (the trend)",
           "All 84 features are physics-readable: a rising s3 slope",
           "means the HPC outlet is actively getting hotter",
           "No black-box embeddings; nothing the domain can't interrogate"])

add_slide("Validation: designed to be hard to fool",
          ["Split by ENGINE, never by window (overlaps would leak)",
           "Model selection on 20 validation engines; test set scored ONCE",
           "Repeated grouped CV: 3 folds × 3 seeds = 9 fits per model",
           "Unit tests for windowing, labels, and metrics (pytest, 10 tests)"],
          img=FIG / "09_cv_stability.png", img_w=Inches(6.5))

add_slide("Results",
          ["HGB: test RMSE 12.74 · PHM 247 · CV 14.40 ± 0.93",
           "RF: 13.48 / 307 · Ridge baseline: 17.43 / 484",
           "Ranking HGB > RF > Ridge holds across all 9 CV fits",
           "Honest expectation for new engines: ~14–15, not 12.7",
           "Most precise near failure: 4.83 RMSE in the 0–30 cycle band"],
          img=FIG / "07_test_rul.png", img_w=Inches(6.5))

add_slide("Explainability: the model watches change, not levels",
          ["7 of top 8 features are _slope — it learned degradation as a process",
           "Top channels: HPC outlet pressure/temp, core speeds",
           "— exactly what HPC-degradation thermodynamics predicts",
           "Ablation: worst single-sensor loss +0.64 RMSE → no sensor is load-bearing",
           "Permutation and ablation disagree on rankings — so neither is trusted alone"],
          img=FIG / "10_permutation_importance.png", img_w=Inches(6.8))

add_slide("Robustness: what's fragile",
          ["Seed-stable: 13.49–14.18 over 5 seeds — not luck",
           "Noise-sensitive: +5.2 RMSE at 0.5× sensor noise — TOP RISK",
           "  (field sensors are noisier than sim sensors → harden with noise-augmented training)",
           "Survives worst-sensor failure via imputation (+1.37) — graceful",
           "Leans optimistic: late errors dominate PHM 2.6× → needs conservative margin"],
          img=FIG / "12_robustness.png", img_w=Inches(6.2))

add_slide("From prediction to product",
          ["Streamlit demo: pick an engine → traces, RUL curve, risk band, action",
           "Decision flow, not just 'RUL = 42': who sees it, what threshold",
           "triggers what, who signs off, what a false alarm costs",
           "Rollout: shadow mode → advisory → constrained automation",
           "KPIs that matter: false alarms/100 cycles, warning lead time,",
           "unplanned removals avoided — not just RMSE"])

add_slide("Limitations — volunteered, each with a next experiment",
          ["Simulated, single condition/fault → FD002–FD004 next",
           "100-point test set is a noisy ruler → wider CV already reported",
           "Point predictions only → quantile intervals for planners",
           "Noise-sensitive → noise-augmented training",
           "Never validated prospectively → shadow-mode protocol",
           "Nothing here claims savings, deployment, or field accuracy."],
          note="The limitations section is a feature: it's what makes the rest believable.")

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
